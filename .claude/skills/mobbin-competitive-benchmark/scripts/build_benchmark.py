#!/usr/bin/env python3
"""Build a competitive-benchmark workbook from a JSON spec.

Usage: python build_benchmark.py spec.json output.xlsx
Then recalculate with LibreOffice (see SKILL.md step 8) so formulas get cached values.
See references/spec-format.md for the JSON structure.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
def font(**k): return Font(name=F, **k)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
_t = Side(style="thin", color="D0D0D0")
BORDER = Border(left=_t, right=_t, top=_t, bottom=_t)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NAVY, STAGE, SCORE, YELLOW, GREY, RED = "1F2A44", "DCE6F2", "D9EAD3", "FFF2CC", "F2F2F2", "F4CCCC"
NC = "Not captured"


def header(ws, row, heads, widths):
    for i, h in enumerate(heads, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        if h == "":
            continue
        c = ws.cell(row, i, h)
        c.font = font(bold=True, color="FFFFFF", size=10); c.fill = fill(NAVY)
        c.alignment = CENTER; c.border = BORDER


def link(cell, url):
    if url:
        cell.hyperlink = url
        cell.font = font(size=9, color="1155CC", underline="single")


def build(spec, out):
    platforms = spec["platforms"]           # list of {name, on_mobbin}
    names = [p["name"] for p in platforms]
    n = len(names)
    title = spec.get("title", "Competitive Benchmark (Mobbin)")
    wb = Workbook()

    # ---- Read Me ----
    rm = wb.active; rm.title = "Read Me"
    rm.column_dimensions["A"].width = 3; rm.column_dimensions["B"].width = 30; rm.column_dimensions["C"].width = 110
    rm["B1"] = title; rm["B1"].font = font(bold=True, size=16, color=NAVY)
    rm["B2"] = spec.get("subtitle", ""); rm["B2"].font = font(italic=True, size=10, color="595959")
    r = 4
    for sec in spec.get("readme", []):
        c = rm.cell(r, 2, sec["heading"]); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(NAVY); rm.cell(r, 3).fill = fill(NAVY)
        r += 1
        for k, v in sec["rows"]:
            a = rm.cell(r, 2, k); a.font = font(bold=True, size=10); a.alignment = WRAP
            b = rm.cell(r, 3, v); b.font = font(size=10); b.alignment = WRAP
            r += 1
        r += 1

    # ---- Benchmark Matrix ----
    mx = wb.create_sheet("Benchmark Matrix")
    heads = ["Funnel stage", "Criterion", "What to look for", "Driver", "Evidence"] + names + ["Best in class", "Implication for our flow"]
    widths = [16, 26, 30, 18, 12] + [30] * n + [18, 34]
    mx["A1"] = title.upper() + " — MATRIX"; mx["A1"].font = font(bold=True, size=14, color=NAVY)
    mx["A2"] = "Observations from Mobbin captures (see Evidence Index). Green rows are editable 1–5 desk scores. Yellow = needs self-testing. Red = notable friction."
    mx["A2"].font = font(italic=True, size=9, color="595959")
    HR = 4; header(mx, HR, heads, widths); mx.row_dimensions[HR].height = 30
    PC0 = 6
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.error = "Enter a whole number from 1 to 5"; mx.add_data_validation(dv)
    score_rows = []; r = HR + 1
    for st in spec["stages"]:
        for col in range(1, len(heads) + 1): mx.cell(r, col).fill = fill(STAGE)
        mx.cell(r, 1, st["name"]).font = font(bold=True, size=11, color=NAVY); r += 1
        short = st["name"].split(". ", 1)[-1].title()
        for cr in st["criteria"]:
            obs = cr["observations"]  # dict platform -> text
            ev = cr.get("evidence", "Mobbin")
            vals = [short, cr["criterion"], cr.get("look_for", ""), cr.get("driver", ""), ev]
            for p in platforms:
                default = NC if p.get("on_mobbin", True) else "Not on Mobbin — self-test"
                vals.append(obs.get(p["name"], default))
            vals += [cr.get("best", ""), cr.get("implication", "")]
            friction = set(cr.get("friction", []))
            for ci, v in enumerate(vals, 1):
                c = mx.cell(r, ci, v); c.font = font(size=9, bold=(ci == 2)); c.alignment = WRAP; c.border = BORDER
                if ci == 5 and "Self-test" in str(v): c.fill = fill(YELLOW)
                if PC0 <= ci < PC0 + n:
                    p = platforms[ci - PC0]
                    if not p.get("on_mobbin", True):
                        c.fill = fill(YELLOW); c.font = font(size=9, italic=True, color="7F6000")
                    elif v == NC:
                        c.fill = fill(GREY); c.font = font(size=9, italic=True, color="808080")
                    elif p["name"] in friction:
                        c.fill = fill(RED)
            mx.row_dimensions[r].height = 78; r += 1
        mx.cell(r, 2, "STAGE SCORE (1–5)").font = font(bold=True, size=10)
        for col in range(1, len(heads) + 1): mx.cell(r, col).fill = fill(SCORE); mx.cell(r, col).border = BORDER
        sc = st.get("scores", {})
        for pi, p in enumerate(platforms):
            v = sc.get(p["name"])
            c = mx.cell(r, PC0 + pi, v); c.alignment = CENTER; c.font = font(bold=True, size=11, color="0000FF")
            if v is None: c.fill = fill(YELLOW)
            dv.add(c)
        c = mx.cell(r, 3, "Desk score from screenshots — edit freely. Yellow = not enough evidence yet.")
        c.font = font(italic=True, size=8, color="38761D"); c.alignment = WRAP
        score_rows.append(r); r += 2
    mx.freeze_panes = mx.cell(HR + 1, 3)

    # ---- Scoring Summary ----
    ss = wb.create_sheet("Scoring Summary")
    ss["B1"] = "SCORING SUMMARY — calculated from the Benchmark Matrix stage-score rows"; ss["B1"].font = font(bold=True, size=14, color=NAVY)
    ss["B2"] = "Formulas — edit scores on the Benchmark Matrix tab. '—' = no score yet."; ss["B2"].font = font(italic=True, size=9, color="595959")
    header(ss, 4, ["", "Funnel stage"] + names, [3, 34] + [15] * n)
    for si, st in enumerate(spec["stages"]):
        rr = 5 + si
        c = ss.cell(rr, 2, st["name"]); c.font = font(size=10); c.border = BORDER
        for pi in range(n):
            col = get_column_letter(PC0 + pi); src = f"'Benchmark Matrix'!{col}{score_rows[si]}"
            c = ss.cell(rr, 3 + pi, f'=IF({src}="","—",{src})'); c.font = font(size=10, color="008000"); c.alignment = CENTER; c.border = BORDER
    avg_r = 5 + len(spec["stages"]); cnt_r = avg_r + 1; rank_r = avg_r + 2
    first, last = get_column_letter(3), get_column_letter(2 + n)
    ss.cell(avg_r, 2, "Average across scored stages").font = font(bold=True, size=10)
    ss.cell(cnt_r, 2, f"Stages scored (of {len(spec['stages'])})").font = font(size=10)
    ss.cell(rank_r, 2, "Rank (by average)").font = font(bold=True, size=10)
    for pi in range(n):
        col = get_column_letter(PC0 + pi); L = get_column_letter(3 + pi)
        rng = ",".join(f"'Benchmark Matrix'!{col}{x}" for x in score_rows)
        a = ss.cell(avg_r, 3 + pi, f'=IFERROR(AVERAGE({rng}),"—")'); a.number_format = "0.0"; a.font = font(bold=True, size=10); a.alignment = CENTER; a.fill = fill(SCORE); a.border = BORDER
        b = ss.cell(cnt_r, 3 + pi, f"=COUNT({rng})"); b.font = font(size=10); b.alignment = CENTER; b.border = BORDER
        k = ss.cell(rank_r, 3 + pi, f'=IF(ISNUMBER({L}{avg_r}),RANK({L}{avg_r},${first}${avg_r}:${last}${avg_r},0),"—")'); k.font = font(bold=True, size=10); k.alignment = CENTER; k.border = BORDER
    ss.cell(rank_r + 2, 2, "Read averages together with 'Stages scored': fewer scored stages means a less comparable average.").font = font(italic=True, size=9, color="595959")
    ch_r = rank_r + 4
    ss.cell(ch_r, 2, "Chart data (0 = not scored)").font = font(size=8, color="808080")
    for pi in range(n):
        L = get_column_letter(3 + pi)
        c = ss.cell(ch_r, 3 + pi, f"=IF(ISNUMBER({L}{avg_r}),{L}{avg_r},0)"); c.font = font(size=8, color="808080"); c.number_format = "0.0"
    ch = BarChart(); ch.type = "col"; ch.title = "Average desk score by platform"; ch.y_axis.title = "Score (1–5)"
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 5; ch.legend = None; ch.height = 8; ch.width = 20
    ch.add_data(Reference(ss, min_col=3, max_col=2 + n, min_row=ch_r), from_rows=True)
    ch.set_categories(Reference(ss, min_col=3, max_col=2 + n, min_row=4))
    ss.add_chart(ch, f"B{ch_r + 2}")

    # ---- Patterns ----
    pt = wb.create_sheet("Patterns")
    pt["A1"] = "PATTERNS — what to copy and what to avoid"; pt["A1"].font = font(bold=True, size=14, color=NAVY)
    header(pt, 3, ["Type", "Stage", "Pattern", "Seen on", "Why it matters", "Mobbin link"], [9, 22, 44, 16, 48, 58])
    colors = {"Copy": "38761D", "Avoid": "CC0000", "Reference": "3D85C6"}
    for i, p in enumerate(spec.get("patterns", [])):
        rr = 4 + i
        for ci, v in enumerate([p["type"], p["stage"], p["pattern"], p["seen_on"], p["why"], p.get("url", "")], 1):
            c = pt.cell(rr, ci, v); c.font = font(size=9); c.alignment = WRAP; c.border = BORDER
            if ci == 1: c.font = font(size=9, bold=True, color="FFFFFF"); c.fill = fill(colors.get(v, "3D85C6"))
            if ci == 6: link(c, v)
        pt.row_dimensions[rr].height = 42
    pt.freeze_panes = "A4"

    # ---- Friction Log ----
    fl = wb.create_sheet("Friction Log")
    fl["A1"] = "FRICTION LOG — observed in Mobbin captures (add your own from self-testing below)"; fl["A1"].font = font(bold=True, size=14, color=NAVY)
    fl["A2"] = "Severity: 3 = likely abandonment, 2 = hesitation or backtracking, 1 = minor annoyance."; fl["A2"].font = font(italic=True, size=9, color="595959")
    header(fl, 4, ["#", "Platform", "Device", "Funnel step", "What happened", "Severity (1-3)", "Evidence", "Design implication for our flow"], [5, 16, 14, 18, 52, 10, 40, 46])
    fr = spec.get("friction", [])
    for i, f in enumerate(fr):
        rr = 5 + i
        vals = [i + 1, f["platform"], f.get("device", "iOS (Mobbin)"), f["step"], f["what"], f["severity"], f.get("url", ""), f["implication"]]
        for ci, v in enumerate(vals, 1):
            c = fl.cell(rr, ci, v); c.font = font(size=9); c.alignment = WRAP; c.border = BORDER
            if ci == 6: c.alignment = CENTER; c.fill = fill({3: "F4CCCC", 2: "FCE5CD", 1: "FFF2CC"}.get(v, YELLOW))
            if ci == 7: link(c, v)
        fl.row_dimensions[rr].height = 40
    start = 5 + len(fr)
    for j in range(10):
        rr = start + j; fl.cell(rr, 1, rr - 4).font = font(size=9)
        for ci in range(2, 9): c = fl.cell(rr, ci); c.fill = fill(YELLOW); c.border = BORDER; c.font = font(size=9)
    fl.freeze_panes = "A5"

    # ---- Evidence Index ----
    ei = wb.create_sheet("Evidence Index")
    ei["A1"] = "EVIDENCE INDEX — Mobbin flows and screens used"; ei["A1"].font = font(bold=True, size=14, color=NAVY)
    ei["A2"] = "Capture context is read from the screenshots (market, currency, dates). Apps may have changed since capture."; ei["A2"].font = font(italic=True, size=9, color="595959")
    header(ei, 4, ["Platform", "Flow (Mobbin name)", "Stages covered", "Screens", "Capture context", "Mobbin link"], [16, 24, 26, 9, 30, 62])
    evid = list(spec.get("evidence", []))
    for p in platforms:
        if not p.get("on_mobbin", True):
            evid.append({"platform": p["name"], "flow": "—", "stages": "—", "screens": "—", "context": "Not found on Mobbin — self-test", "url": ""})
    for i, e in enumerate(evid):
        rr = 5 + i
        for ci, v in enumerate([e["platform"], e["flow"], e["stages"], e["screens"], e["context"], e.get("url", "")], 1):
            c = ei.cell(rr, ci, v); c.font = font(size=9); c.alignment = WRAP; c.border = BORDER
            if ci == 6: link(c, v)
            if e["context"].startswith("Not found"): c.fill = fill(YELLOW)
    ei.freeze_panes = "A5"

    wb.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: build_benchmark.py spec.json output.xlsx")
    with open(sys.argv[1]) as f:
        print(build(json.load(f), sys.argv[2]))
